from pathlib import Path
import json,re,sys,runpy
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(__file__).parent;OUT=ROOT.parent/'dist/assets/documents'
build_journey=runpy.run_path(str(ROOT/'build-journey.py'))['build']
D=json.loads((ROOT/'cv-content.json').read_text(encoding='utf8'));A=json.loads((ROOT/'author-content.json').read_text(encoding='utf8'))
R=json.loads((ROOT/'revised-content.json').read_text(encoding='utf8'))
for l in ['uk','en','de','pl']:R[l]=json.loads((ROOT/f'translations-{l}.json').read_text(encoding='utf8'))
L={
'ru':['CV художницы','Творческий путь','Выставки и проекты','Продолжение выставочного списка','Резиденции','Награды и стипендии','Образование','Практический опыт','Языки','Профессиональное членство','Белые маки 2022','Авторская керамическая композиция'],
'uk':['CV художниці','Творчий шлях','Виставки та проєкти','Продовження виставкового списку','Резиденції','Нагороди та стипендії','Освіта','Практичний досвід','Мови','Професійне членство','Білі маки 2022','Авторська керамічна композиція'],
'en':['Artist CV','Artistic journey','Exhibitions and projects','Exhibitions continued','Residencies','Awards and scholarships','Education','Professional practice','Languages','Professional membership','White Poppies 2022','Original ceramic composition'],
'de':['Künstlerischer Lebenslauf','Künstlerischer Werdegang','Ausstellungen und Projekte','Ausstellungen Fortsetzung','Residenzen','Auszeichnungen und Stipendien','Ausbildung','Berufliche Praxis','Sprachkenntnisse','Mitgliedschaft','Weißer Mohn 2022','Künstlerische Keramikkomposition'],
'pl':['CV artystyczne','Droga artystyczna','Wystawy i projekty','Wystawy ciąg dalszy','Rezydencje','Nagrody i stypendia','Wykształcenie','Praktyka zawodowa','Języki','Członkostwo','Białe maki 2022','Autorska kompozycja ceramiczna']}
INK='252420';MUTED='666158';PAPER='F2EEE6';BROWN='574B36'
def text(s):return re.sub(r'\s+',' ',s).strip()
def shade(cell,color):
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(sh)
def init(title,right=False):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.6);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
 for name,size in [('Normal',10.5),('Title',29),('Heading 1',11),('Heading 2',12),('Caption',9)]:
  st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(INK);st.paragraph_format.alignment=3 if right else 0;st.paragraph_format.line_spacing=1.12;st.paragraph_format.space_after=Pt(6)
  if name!='Normal':st.paragraph_format.keep_with_next=True
 d.styles['Title'].paragraph_format.alignment=0;d.styles['Heading 1'].paragraph_format.alignment=0;d.styles['Title'].font.bold=True;d.styles['Title'].font.color.rgb=RGBColor(0,0,0);d.styles['Title'].paragraph_format.space_after=Pt(7)
 d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(13);d.styles['Heading 1'].paragraph_format.space_after=Pt(7)
 d.styles['Heading 2'].font.bold=True;d.styles['Heading 2'].paragraph_format.space_before=Pt(10)
 for n in list(d.styles.element.iter(qn('w:pBdr'))):n.getparent().remove(n)
 d.core_properties.author='Krystyna Mishchenko';d.core_properties.title=title
 f=s.footer.paragraphs[0];f.alignment=2;run=f.add_run('KRYSTYNA MISHCHENKO   /   ');run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string(MUTED);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 return d
def p(d,s,style=None):return d.add_paragraph(text(s),style)
def head(d,s):return p(d,s.upper(),'Heading 1')
def header(d,l,kind,continued=False):
 a=A[l];r=p(d,L[l][kind].upper());r.paragraph_format.space_after=Pt(11)
 for run in r.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string(BROWN)
 name=p(d,D[l]['name'].upper(),'Title')
 for run in name.runs:
  spacing=OxmlElement('w:spacing');spacing.set(qn('w:val'),'24');run._r.get_or_add_rPr().append(spacing)
 p(d,a['role'])
 band=d.add_table(rows=1,cols=1);cell=band.cell(0,0);shade(cell,PAPER)
 cp=cell.paragraphs[0];cp.add_run(a['born']+'\n'+D[l]['city']);cp.paragraph_format.space_after=Pt(3)
 cp=cell.add_paragraph('m.krystyna.197@gmail.com  ·  DE +49 1575 3604617  ·  UA +380 93 1324782');cp.paragraph_format.space_after=Pt(3)
 for cp in cell.paragraphs:
  cp.alignment=0
  for run in cp.runs:run.font.size=Pt(8.5)
 spacer=d.add_paragraph();spacer.paragraph_format.space_after=Pt(8);spacer.paragraph_format.line_spacing=0.3
 return d
def chronology(d,items):
 table=d.add_table(rows=0,cols=2);table.autofit=False;table.columns[0].width=Cm(3.2);table.columns[1].width=Cm(14.2)
 for item in items:
  date,sep,body=item.partition(' | ')
  if not sep:date,body='',item
  cells=table.add_row().cells;cells[0].width=Cm(3.2);cells[1].width=Cm(14.2)
  for i,s in enumerate([date,body]):
   cp=cells[i].paragraphs[0];cp.alignment=2 if i==0 else 3;cp.paragraph_format.space_after=Pt(5);cp.paragraph_format.line_spacing=1.06;r=cp.add_run(text(s));r.font.size=Pt(9.5 if i==0 else 10)
   if i==0:r.font.color.rgb=RGBColor.from_string(MUTED);r.bold=True
  no=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(no)
 return table
def language_grid(doc,items):
 table=doc.add_table(rows=0,cols=2)
 for i in range(0,len(items),2):
  cells=table.add_row().cells
  for cell,value in zip(cells,items[i:i+2]):
   cp=cell.paragraphs[0];cp.alignment=0;cp.add_run(value.strip())

for l,d in D.items():
 if sys.argv[1:]and l not in sys.argv[1:]:continue
 a=A[l];lab=L[l];r=R[l]
 doc=header(init(lab[0],True),l,0);p(doc,a['profile']);head(doc,lab[6])
 edu=[]
 for i in range(3):
  title=d['edu'][i];date=re.search(r'\d{4}–\d{4}',title).group();qualification=title.replace(date,'').strip(' ,');school=d['edu'][3]if i<2 else d['edu'][5];edu.append(date+' | '+qualification+' · '+school)
 chronology(doc,edu)
 head(doc,lab[9]);p(doc,d['member'])
 head(doc,lab[7]);chronology(doc,[date+' | '+studio+' · '+r['jobs'][i]for i,(studio,date)in enumerate([('SAGA POTTERY','11.2020–02.2022'),('PROCS','2018–2021'),('Studio Gozhar','2020')])]);head(doc,lab[8]);language_grid(doc,d['langs'].split(';'))
 doc.add_page_break();p(doc,d['name'].upper(),'Title')
 if d.get('residencies'):head(doc,lab[4]);chronology(doc,d['residencies'])
 head(doc,lab[2]);chronology(doc,d['ex']);head(doc,lab[5]);chronology(doc,d.get('scholarships',[])+d['aw'])
 doc.save(OUT/f'Krystyna-Mishchenko-CV-{l.upper()}.docx')
 build_journey([l]);print(l,'updated CV and approved artistic journey')
site={l:{'name':d['name'],'city':d['city'],'birthYear':1997,'education':d['edu'],'member':d['member'],'head':d['head'],'profile':A[l]['profile'],'author':A[l],'jobs':R[l]['jobs'],'langs':d['langs'].split(';'),'ex':d['ex'],'aw':d['aw'],'residencies':d.get('residencies',[]),'scholarships':d.get('scholarships',[])}for l,d in D.items()}
(ROOT.parent/'dist/cv-data.js').write_text('window.portfolioCV='+json.dumps(site,ensure_ascii=False)+';\n',encoding='utf8')
