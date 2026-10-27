"""Publish the artist-approved Russian manuscript and faithful translations."""
from pathlib import Path
import json, shutil
from docx import Document

ROOT = Path(__file__).parent
OUT = ROOT.parent / 'dist/assets/documents'

def replace(paragraph, value):
    # Retain the approved paragraph layout and first run's character styling.
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(value)

def build(languages=None):
    template = ROOT / 'approved-artistic-journey-ru.docx'
    translations = json.loads((ROOT / 'journey-translations.json').read_text(encoding='utf-8'))
    author = json.loads((ROOT / 'author-content.json').read_text(encoding='utf-8'))
    cv = json.loads((ROOT / 'cv-content.json').read_text(encoding='utf-8'))
    headings = {'uk':'ТВОРЧИЙ ШЛЯХ','en':'ARTISTIC JOURNEY','de':'KÜNSTLERISCHER WERDEGANG','pl':'DROGA ARTYSTYCZNA'}
    for lang in languages or ['ru','uk','en','de','pl']:
        target = OUT / f'Krystyna-Mishchenko-Artistic-Journey-{lang.upper()}.docx'
        if lang == 'ru':
            shutil.copyfile(template, target)
            continue
        doc = Document(template)
        assert len(translations[lang]) == len(doc.paragraphs[4:]) == 10
        replace(doc.paragraphs[0], headings[lang])
        replace(doc.paragraphs[1], cv[lang]['name'].upper())
        replace(doc.paragraphs[2], author[lang]['role'])
        for paragraph, value in zip(doc.paragraphs[4:], translations[lang]):
            replace(paragraph, value)
        cell = doc.tables[0].cell(0, 0)
        replace(cell.paragraphs[0], author[lang]['born'] + '\n' + cv[lang]['city'])
        doc.core_properties.title = headings[lang].capitalize()
        doc.save(target)

if __name__ == '__main__':
    build()
